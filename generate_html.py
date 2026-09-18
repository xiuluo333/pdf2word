"""Update vocabulary metadata and generate a browser-based reader.

Run ``python generate_html.py`` to add unique IDs and local ``words.json`` phonetics to
``考研生词.xlsx``, then create ``考研生词.html``.  The generated document is
self-contained: workbook data is embedded as JSON and all interaction is
implemented in the page itself, so it can be opened directly from disk.
"""

from __future__ import annotations

import argparse
import json
import os
import re
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree


BASE_DIR = Path(__file__).resolve().parent


FIELD_ALIASES = {
    "id": ("id", "word id", "word_id", "单词编号", "唯一编号"),
    "word": ("word", "单词", "词汇"),
    "translate": ("translate", "translation", "释义", "翻译", "中文释义"),
    "english_definition": ("English definition", "english_definition", "英译英", "英文释义"),
    "number": ("number", "编号", "频次", "frequency"),
    "part_of_speech": ("part of speech", "part_of_speech", "词性"),
    "transformation": ("transformation", "word forms", "词形变化"),
    "memory_techniques": ("memory techniques", "memory_technique", "记忆方法"),
    "similar_words": (
        "distinguishing between similar words",
        "similar words",
        "近义词辨析",
        "易混淆词",
    ),
    "phonetic": (
        "phonetic",
        "phonetics",
        "ipa",
        "音标",
        "英式音标",
        "美式音标",
    ),
    "example_sentence": (
        "example sentence", "example", "sentence", "例句", "日常例句", "对话例句"
    ),
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Add vocabulary IDs/phonetics to Excel and generate its study webpage.")
    parser.add_argument("--input", type=Path, default=Path(os.getenv("VOCAB_INPUT", str(BASE_DIR / "考研生词.xlsx"))))
    parser.add_argument("--output", type=Path, default=None)
    parser.add_argument("--sheet", default=None, help="Worksheet name; defaults to the active worksheet.")
    parser.add_argument("--words", type=Path, default=None, help="Phonetic JSON path; defaults to words.json beside the workbook.")
    return parser.parse_args()


def normalise(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).lower()


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def find_columns(headers: list[Any]) -> dict[str, int | None]:
    by_name = {normalise(value): i for i, value in enumerate(headers) if value is not None}
    result: dict[str, int | None] = {}
    for field, aliases in FIELD_ALIASES.items():
        result[field] = next((by_name[normalise(alias)] for alias in aliases if normalise(alias) in by_name), None)
    if result["word"] is None:
        raise ValueError("Workbook must contain a 'word' (or 单词) column")
    return result


def load_word_data(path: Path) -> dict[str, dict[str, str]]:
    """Load word metadata (phonetic and dictionary meaning) from words.json."""

    if not path.exists():
        return {}
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"Could not read phonetic JSON: {path}") from exc

    entries: Any
    if isinstance(payload, list):
        entries = payload
    elif isinstance(payload, dict):
        # Also accept a compact {"word": "/.../"} mapping.
        entries = [
            {"word": word, "phonetic": phonetic}
            for word, phonetic in payload.items()
            if isinstance(word, str)
        ]
    else:
        raise ValueError("Phonetic JSON must be a list or an object mapping words to phonetics")

    result: dict[str, dict[str, str]] = {}
    for entry in entries:
        if not isinstance(entry, dict):
            continue
        word = cell_text(entry.get("word"))
        if word:
            result.setdefault(word.casefold(), {
                "phonetic": cell_text(entry.get("phonetic")),
                "translate": cell_text(entry.get("meaning") or entry.get("translate") or entry.get("translation")),
            })
    return result


def load_phonetics(path: Path) -> dict[str, str]:
    """Backward-compatible helper returning only phonetics."""
    return {word: data["phonetic"] for word, data in load_word_data(path).items() if data.get("phonetic")}


def next_word_id(used_ids: set[str]) -> str:
    """Return a readable, stable-looking ID that cannot collide in the sheet."""

    numbers = []
    for value in used_ids:
        match = re.fullmatch(r"W(\d+)", value, flags=re.IGNORECASE)
        if match:
            numbers.append(int(match.group(1)))
    candidate = max(numbers, default=0) + 1
    while f"W{candidate:06d}" in used_ids:
        candidate += 1
    return f"W{candidate:06d}"


def _rows_from_values(headers: list[Any], data_rows: list[tuple[int, list[Any]]]) -> list[dict[str, str]]:
    """Convert worksheet values into the small, stable JSON schema used by the page."""

    columns = find_columns(headers)
    fields = ("translate", "english_definition", "number", "part_of_speech", "transformation", "memory_techniques", "similar_words", "phonetic", "example_sentence")
    rows: list[dict[str, str]] = []
    for source_index, values in data_rows:
        word_column = columns["word"]
        word = cell_text(values[word_column]) if word_column is not None and word_column < len(values) else ""
        if not word:
            continue
        id_column = columns["id"]
        row_id = cell_text(values[id_column]) if id_column is not None and id_column < len(values) else ""
        row = {"id": row_id or f"W{len(rows) + 1:06d}", "sourceRow": str(source_index), "word": word}
        for field in fields:
            column = columns[field]
            row[field] = cell_text(values[column]) if column is not None and column < len(values) else ""
        rows.append(row)
    return rows


def _load_rows_with_xml(path: Path, sheet_name: str | None = None) -> list[dict[str, str]]:
    """Read normal XLSX cell types without third-party packages.

    This deliberately covers the worksheet features needed for a vocabulary
    table (inline strings, shared strings, numbers, booleans and blanks).  If
    openpyxl is available it remains the preferred parser.
    """

    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ns = {"m": main_ns, "r": rel_ns, "pr": package_rel_ns}
    with ZipFile(path) as archive:
        workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        sheets = workbook.find("m:sheets", ns)
        if sheets is None:
            raise ValueError("Workbook does not contain any worksheets")
        sheet_nodes = list(sheets)
        selected = None
        if sheet_name:
            selected = next((node for node in sheet_nodes if node.attrib.get("name") == sheet_name), None)
            if selected is None:
                raise ValueError(f"Worksheet not found: {sheet_name}")
        else:
            active_tab = 0
            views = workbook.find("m:bookViews", ns)
            if views is not None:
                view = views.find("m:workbookView", ns)
                if view is not None:
                    try:
                        active_tab = max(0, int(view.attrib.get("activeTab", "0")))
                    except ValueError:
                        active_tab = 0
            selected = sheet_nodes[min(active_tab, len(sheet_nodes) - 1)]

        relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        relationship_map = {node.attrib.get("Id"): node.attrib.get("Target", "") for node in relationships}
        target = relationship_map.get(selected.attrib.get(f"{{{rel_ns}}}id"), "")
        target = target.lstrip("/")
        if not target.startswith("xl/"):
            target = f"xl/{target}"
        sheet_root = ElementTree.fromstring(archive.read(target))

        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in shared_root.findall("m:si", ns):
                shared_strings.append("".join(item.itertext()))

        def read_cell(cell: ElementTree.Element) -> str:
            value = cell.find("m:v", ns)
            cell_type = cell.attrib.get("t", "")
            if cell_type == "inlineStr":
                inline = cell.find("m:is", ns)
                return "".join(inline.itertext()) if inline is not None else ""
            if value is None or value.text is None:
                return ""
            raw = value.text
            if cell_type == "s":
                try:
                    return shared_strings[int(raw)]
                except (ValueError, IndexError):
                    return ""
            if cell_type == "b":
                return "TRUE" if raw == "1" else "FALSE"
            return raw

        def column_index(reference: str) -> int:
            letters = re.match(r"[A-Za-z]+", reference or "")
            if not letters:
                return 0
            index = 0
            for letter in letters.group(0).upper():
                index = index * 26 + ord(letter) - ord("A") + 1
            return index - 1

        sheet_data = sheet_root.find("m:sheetData", ns)
        if sheet_data is None:
            raise ValueError("Worksheet is empty")
        parsed_rows: list[tuple[int, dict[int, str]]] = []
        for row_node in sheet_data.findall("m:row", ns):
            row_number = int(row_node.attrib.get("r", len(parsed_rows) + 1))
            cells = {column_index(cell.attrib.get("r", "")): read_cell(cell) for cell in row_node.findall("m:c", ns)}
            parsed_rows.append((row_number, cells))
        if not parsed_rows:
            raise ValueError("Worksheet is empty")
        header_map = parsed_rows[0][1]
        max_header_column = max(header_map, default=-1)
        headers = [header_map.get(index, "") for index in range(max_header_column + 1)]
        data_rows = []
        for row_number, values in parsed_rows[1:]:
            width = max(max_header_column + 1, max(values, default=-1) + 1)
            data_rows.append((row_number, [values.get(index, "") for index in range(width)]))
    return _rows_from_values(headers, data_rows)


def load_rows(path: Path, sheet_name: str | None = None) -> list[dict[str, str]]:
    try:
        import openpyxl
    except ImportError:
        return _load_rows_with_xml(path, sheet_name)

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = list(next(iterator))
        except StopIteration as exc:
            raise ValueError("Workbook is empty") from exc
        return _rows_from_values(headers, list(enumerate(iterator, start=2)))
    finally:
        workbook.close()


def _update_workbook_with_openpyxl(path: Path, word_data: dict[str, dict[str, str]], sheet_name: str | None) -> None:
    import openpyxl

    workbook = openpyxl.load_workbook(path)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        headers = [cell.value for cell in sheet[1]]
        columns = find_columns(headers)
        word_column = columns["word"]
        assert word_column is not None
        next_column = sheet.max_column
        id_column = columns["id"]
        if id_column is None:
            id_column = next_column
            next_column += 1
        phonetic_column = columns["phonetic"]
        if phonetic_column is None:
            phonetic_column = next_column
            next_column += 1
        translate_column = columns["translate"]
        if translate_column is None:
            translate_column = next_column
            next_column += 1
            sheet.cell(1, translate_column + 1).value = "translate"
        if columns["id"] is None:
            sheet.cell(1, id_column + 1).value = "id"
        if columns["phonetic"] is None:
            sheet.cell(1, phonetic_column + 1).value = "phonetic"

        used_ids: set[str] = set()
        for row_number in range(2, sheet.max_row + 1):
            word = cell_text(sheet.cell(row_number, word_column + 1).value)
            if not word:
                continue
            id_cell = sheet.cell(row_number, id_column + 1)
            current_id = cell_text(id_cell.value)
            if not current_id or current_id in used_ids:
                current_id = next_word_id(used_ids)
                id_cell.value = current_id
            used_ids.add(current_id)

            phonetic_cell = sheet.cell(row_number, phonetic_column + 1)
            if not cell_text(phonetic_cell.value):
                metadata = word_data.get(word.casefold(), {})
                phonetic = metadata.get("phonetic", "")
                if phonetic:
                    phonetic_cell.value = phonetic
            translate_cell = sheet.cell(row_number, translate_column + 1)
            translation = word_data.get(word.casefold(), {}).get("translate", "")
            if translation:
                # The local dictionary is the authoritative source for this field.
                translate_cell.value = translation
        workbook.save(path)
    finally:
        workbook.close()


def _xml_column_letters(index: int) -> str:
    result = ""
    value = index + 1
    while value:
        value, remainder = divmod(value - 1, 26)
        result = chr(ord("A") + remainder) + result
    return result


def _update_workbook_with_xml(path: Path, word_data: dict[str, dict[str, str]], sheet_name: str | None) -> None:
    """Update metadata in an XLSX when openpyxl is unavailable."""

    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ns = {"m": main_ns, "r": rel_ns, "pr": package_rel_ns}
    temp_path = path.with_name(path.name + ".tmp")
    try:
        with ZipFile(path) as archive:
            workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            sheets = workbook.find("m:sheets", ns)
            if sheets is None or not list(sheets):
                raise ValueError("Workbook does not contain any worksheets")
            sheet_nodes = list(sheets)
            if sheet_name:
                selected = next((node for node in sheet_nodes if node.attrib.get("name") == sheet_name), None)
                if selected is None:
                    raise ValueError(f"Worksheet not found: {sheet_name}")
            else:
                active_tab = 0
                views = workbook.find("m:bookViews", ns)
                if views is not None:
                    view = views.find("m:workbookView", ns)
                    if view is not None:
                        try:
                            active_tab = max(0, int(view.attrib.get("activeTab", "0")))
                        except ValueError:
                            active_tab = 0
                selected = sheet_nodes[min(active_tab, len(sheet_nodes) - 1)]

            relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            relationship_map = {node.attrib.get("Id"): node.attrib.get("Target", "") for node in relationships}
            target = relationship_map.get(selected.attrib.get(f"{{{rel_ns}}}id"), "").lstrip("/")
            if not target.startswith("xl/"):
                target = f"xl/{target}"
            sheet_root = ElementTree.fromstring(archive.read(target))
            sheet_data = sheet_root.find("m:sheetData", ns)
            if sheet_data is None:
                raise ValueError("Worksheet is empty")
            sheet_rows = list(sheet_data.findall("m:row", ns))
            if not sheet_rows:
                raise ValueError("Worksheet is empty")

            shared_strings: list[str] = []
            if "xl/sharedStrings.xml" in archive.namelist():
                shared_root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
                shared_strings = ["".join(item.itertext()) for item in shared_root.findall("m:si", ns)]

            def column_index(reference: str) -> int:
                letters = re.match(r"[A-Za-z]+", reference or "")
                if not letters:
                    return 0
                index = 0
                for letter in letters.group(0).upper():
                    index = index * 26 + ord(letter) - ord("A") + 1
                return index - 1

            def read_cell(cell: ElementTree.Element) -> str:
                value = cell.find("m:v", ns)
                cell_type = cell.attrib.get("t", "")
                if cell_type == "inlineStr":
                    inline = cell.find("m:is", ns)
                    return "".join(inline.itertext()) if inline is not None else ""
                if value is None or value.text is None:
                    return ""
                if cell_type == "s":
                    try:
                        return shared_strings[int(value.text)]
                    except (ValueError, IndexError):
                        return ""
                return value.text

            def row_number(row: ElementTree.Element, fallback: int) -> int:
                try:
                    return int(row.attrib.get("r", fallback))
                except ValueError:
                    return fallback

            header_row = sheet_rows[0]
            header_cells = {column_index(cell.attrib.get("r", "")): read_cell(cell) for cell in header_row.findall("m:c", ns)}
            max_column = max((column_index(cell.attrib.get("r", "")) for row in sheet_rows for cell in row.findall("m:c", ns)), default=-1)
            headers = [header_cells.get(index, "") for index in range(max(max_column, max(header_cells, default=-1)) + 1)]
            columns = find_columns(headers)
            word_column = columns["word"]
            assert word_column is not None
            id_column = columns["id"]
            if id_column is None:
                max_column += 1
                id_column = max_column
            phonetic_column = columns["phonetic"]
            if phonetic_column is None:
                max_column += 1
                phonetic_column = max_column
            translate_column = columns["translate"]
            if translate_column is None:
                max_column += 1
                translate_column = max_column

            def find_or_create_cell(row: ElementTree.Element, column: int, number: int) -> ElementTree.Element:
                for cell in row.findall("m:c", ns):
                    if column_index(cell.attrib.get("r", "")) == column:
                        return cell
                cell = ElementTree.Element(f"{{{main_ns}}}c", {"r": f"{_xml_column_letters(column)}{number}"})
                children = row.findall("m:c", ns)
                insert_at = len(children)
                for index, existing in enumerate(children):
                    if column_index(existing.attrib.get("r", "")) > column:
                        insert_at = index
                        break
                row.insert(insert_at, cell)
                return cell

            def set_inline_string(cell: ElementTree.Element, value: str) -> None:
                cell.attrib["t"] = "inlineStr"
                for child in list(cell):
                    cell.remove(child)
                inline = ElementTree.SubElement(cell, f"{{{main_ns}}}is")
                text = ElementTree.SubElement(inline, f"{{{main_ns}}}t")
                text.text = value

            first_number = row_number(header_row, 1)
            set_inline_string(find_or_create_cell(header_row, id_column, first_number), "id")
            set_inline_string(find_or_create_cell(header_row, phonetic_column, first_number), "phonetic")
            set_inline_string(find_or_create_cell(header_row, translate_column, first_number), "translate")
            used_ids: set[str] = set()
            for position, row in enumerate(sheet_rows[1:], start=2):
                number = row_number(row, position)
                cells = {column_index(cell.attrib.get("r", "")): cell for cell in row.findall("m:c", ns)}
                word = read_cell(cells[word_column]) if word_column in cells else ""
                if not word:
                    continue
                id_cell = find_or_create_cell(row, id_column, number)
                current_id = read_cell(id_cell)
                if not current_id or current_id in used_ids:
                    current_id = next_word_id(used_ids)
                    set_inline_string(id_cell, current_id)
                used_ids.add(current_id)
                phonetic_cell = find_or_create_cell(row, phonetic_column, number)
                if not read_cell(phonetic_cell):
                    phonetic = word_data.get(word.casefold(), {}).get("phonetic", "")
                    if phonetic:
                        set_inline_string(phonetic_cell, phonetic)
                translate_cell = find_or_create_cell(row, translate_column, number)
                translation = word_data.get(word.casefold(), {}).get("translate", "")
                if translation:
                    set_inline_string(translate_cell, translation)

            dimension = sheet_root.find("m:dimension", ns)
            if dimension is not None:
                last_row = max(row_number(row, index + 1) for index, row in enumerate(sheet_rows))
                dimension.attrib["ref"] = f"A1:{_xml_column_letters(max_column)}{last_row}"
            ElementTree.register_namespace("", main_ns)
            ElementTree.register_namespace("r", rel_ns)
            updated_sheet = ElementTree.tostring(sheet_root, encoding="utf-8", xml_declaration=True)
            with ZipFile(temp_path, "w", ZIP_DEFLATED) as output:
                for item in archive.infolist():
                    output.writestr(item, updated_sheet if item.filename == target else archive.read(item.filename))
        os.replace(temp_path, path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def update_workbook(path: Path, words_path: Path, sheet_name: str | None = None) -> None:
    word_data = load_word_data(words_path)
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        _update_workbook_with_xml(path, word_data, sheet_name)
    else:
        _update_workbook_with_openpyxl(path, word_data, sheet_name)


HTML_TEMPLATE = r'''<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <meta name="theme-color" content="#18252b">
  <title>考研生词 · 学习卡片</title>
  <style>
    :root { color-scheme: light; --ink: #18252b; --muted: #68757a; --line: #dce5e5; --paper: #fff; --wash: #f4f7f6; --accent: #0c766d; --accent-soft: #e5f3f0; --gold: #c9872c; }
    * { box-sizing: border-box; }
    html { scroll-behavior: smooth; }
    body { margin: 0; min-width: 320px; background: var(--wash); color: var(--ink); font-family: "Noto Sans SC", "PingFang SC", "Microsoft YaHei", system-ui, sans-serif; }
    button, input { font: inherit; }
    button { cursor: pointer; }
    .topbar { position: sticky; z-index: 10; top: 0; background: rgba(244,247,246,.94); border-bottom: 1px solid var(--line); backdrop-filter: blur(16px); }
    .topbar-inner { max-width: 1360px; margin: 0 auto; padding: 22px 30px 18px; }
    .masthead { display: flex; align-items: flex-start; justify-content: space-between; gap: 20px; }
    h1 { margin: 0; font-size: clamp(1.55rem, 3vw, 2.45rem); letter-spacing: .01em; line-height: 1.1; }
    .kicker { margin: 0 0 8px; color: var(--accent); font-size: .74rem; font-weight: 800; letter-spacing: .13em; text-transform: uppercase; }
    .subtitle { margin: 9px 0 0; color: var(--muted); font-size: .9rem; }
    .resume-button { display: none; border: 1px solid var(--accent); border-radius: 5px; padding: 9px 13px; color: var(--accent); background: var(--paper); font-weight: 700; white-space: nowrap; }
    .resume-button.visible { display: inline-flex; align-items: center; gap: 7px; }
    .resume-button:hover { background: var(--accent-soft); }
    .toolbar { display: grid; grid-template-columns: minmax(200px, 1fr) auto auto; align-items: center; gap: 15px; margin-top: 22px; }
    .search-wrap { position: relative; }
    .search-wrap svg { position: absolute; left: 15px; top: 50%; width: 18px; height: 18px; color: var(--muted); transform: translateY(-50%); pointer-events: none; }
    #search { width: 100%; border: 1px solid #cbd8d7; border-radius: 5px; padding: 13px 15px 13px 44px; outline: 0; background: var(--paper); color: var(--ink); }
    #search:focus { border-color: var(--accent); box-shadow: 0 0 0 3px rgba(12,118,109,.12); }
    .voice-control { display: inline-flex; align-items: center; gap: 8px; border: 1px solid #cbd8d7; border-radius: 5px; padding: 8px 10px; background: var(--paper); color: var(--muted); font-size: .78rem; white-space: nowrap; }
    #voice-select { border: 0; outline: 0; background: transparent; color: var(--ink); font-weight: 700; cursor: pointer; }
    #voice-select:focus-visible { box-shadow: 0 0 0 3px rgba(12,118,109,.12); }
    .stats { display: flex; align-items: center; justify-content: flex-end; gap: 16px; color: var(--muted); font-size: .82rem; white-space: nowrap; }
    .stats strong { color: var(--ink); font-size: 1rem; }
    .meter { width: 110px; height: 6px; overflow: hidden; border-radius: 9px; background: #d8e4e2; }
    .meter i { display: block; width: 0; height: 100%; border-radius: inherit; background: var(--accent); transition: width .25s ease; }
    .search-results { position: absolute; z-index: 11; top: calc(100% + 6px); left: 0; right: 0; display: none; overflow: hidden; border: 1px solid var(--line); border-radius: 5px; background: var(--paper); box-shadow: 0 12px 30px rgba(24,37,43,.12); }
    .search-results.open { display: block; }
    .search-result { display: flex; width: 100%; align-items: baseline; justify-content: space-between; gap: 12px; border: 0; border-bottom: 1px solid #edf2f1; padding: 11px 14px; text-align: left; background: transparent; color: var(--ink); }
    .search-result:last-child { border-bottom: 0; }
    .search-result:hover, .search-result:focus-visible { background: var(--accent-soft); outline: 0; }
    .search-result small { color: var(--muted); }
    main { max-width: 1360px; margin: 0 auto; padding: 28px 30px 70px; }
    .section-line { display: flex; align-items: center; justify-content: space-between; margin: 0 0 17px; color: var(--muted); font-size: .78rem; }
    .section-line span:first-child { font-weight: 700; letter-spacing: .08em; text-transform: uppercase; }
    /* Grid auto-flow is intentionally row-major: frequency order reads left-to-right. */
    #grid { display: grid; grid-template-columns: minmax(0, 1fr); align-items: start; gap: 18px; }
    @media (min-width: 650px) { #grid { grid-template-columns: repeat(2, minmax(0, 1fr)); } }
    @media (min-width: 980px) { #grid { grid-template-columns: repeat(3, minmax(0, 1fr)); } }
    @media (min-width: 1320px) { #grid { grid-template-columns: repeat(4, minmax(0, 1fr)); } }
    .card { display: block; width: 100%; margin: 0; overflow: hidden; border: 1px solid var(--line); border-radius: 7px; background: var(--paper); box-shadow: 0 2px 4px rgba(24,37,43,.025); transition: border-color .2s, box-shadow .2s, transform .2s; scroll-margin-top: 148px; }
    .card:hover { border-color: #b8cfcb; box-shadow: 0 8px 22px rgba(24,37,43,.08); transform: translateY(-1px); }
    .card.current { border-color: var(--accent); box-shadow: 0 0 0 3px rgba(12,118,109,.10); }
    .card-head { display: flex; align-items: flex-start; justify-content: space-between; gap: 12px; padding: 18px 18px 13px; }
    .word { margin: 0; font-family: Georgia, "Times New Roman", serif; font-size: 1.75rem; letter-spacing: .01em; line-height: 1; cursor: help; }
    .word:focus-visible { outline: 2px solid var(--accent); outline-offset: 4px; border-radius: 2px; }
    .word-meta { display: flex; align-items: center; gap: 7px; margin-top: 8px; color: var(--muted); font-size: .82rem; }
    .phonetic { color: var(--accent); font-family: Georgia, serif; font-size: .95rem; }
    .phonetic.missing { color: #9ca9a7; font-family: inherit; font-size: .76rem; }
    .speak { display: inline-grid; flex: 0 0 auto; place-items: center; width: 34px; height: 34px; border: 1px solid #cbd8d7; border-radius: 50%; background: transparent; color: var(--accent); }
    .speak:hover, .speak:focus-visible { border-color: var(--accent); background: var(--accent-soft); outline: 0; }
    .speak svg { width: 16px; height: 16px; }
    .phonetic-button { display: inline-flex; align-items: center; gap: 4px; border: 0; padding: 0; color: inherit; background: transparent; }
    .phonetic-button:hover .phonetic { text-decoration: underline; text-underline-offset: 3px; }
    .badges { display: flex; flex-direction: column; align-items: flex-end; gap: 7px; }
    .word-id { color: var(--muted); font-family: ui-monospace, SFMono-Regular, Menlo, monospace; font-size: .68rem; letter-spacing: .04em; }
    .number { flex: 0 0 auto; border-radius: 4px; padding: 4px 7px; color: var(--gold); background: #fff7e9; font-size: .72rem; font-weight: 800; }
    .meaning { margin: 0 18px 15px; border-left: 3px solid var(--accent); padding: 3px 0 3px 11px; color: #263b3e; font-size: .93rem; line-height: 1.65; }
    .details { border-top: 1px solid #edf2f1; padding: 3px 18px 16px; }
    .detail { display: grid; grid-template-columns: 93px minmax(0, 1fr); gap: 9px; padding-top: 12px; font-size: .82rem; line-height: 1.65; }
    .detail-label { color: var(--muted); font-size: .73rem; font-weight: 800; letter-spacing: .05em; }
    .detail-value { min-width: 0; white-space: pre-wrap; overflow-wrap: anywhere; color: #405154; }
    .detail-value.empty { color: #a4afad; font-style: italic; }
    .example-button { display: inline; border: 0; border-bottom: 1px dashed var(--accent); padding: 0 2px; text-align: left; background: transparent; color: inherit; cursor: pointer; }
    .example-button:hover, .example-button:focus-visible { border-bottom-style: solid; color: var(--accent); outline: 0; }
    .example-button svg { width: 14px; height: 14px; margin-right: 5px; vertical-align: -2px; }
    .empty-state { display: none; padding: 54px 20px; border: 1px dashed #bdcecb; border-radius: 7px; text-align: center; color: var(--muted); }
    .empty-state.visible { display: block; }
    .loading { padding: 20px; color: var(--muted); text-align: center; font-size: .8rem; }
    [hidden] { display: none !important; }
    .study-toolbar, .pagination, .card-actions { display: flex; flex-wrap: wrap; align-items: center; gap: 12px; }
    .study-toolbar { margin-top: 14px; font-size: .8rem; }
    .study-toolbar button, .pagination button, .card-actions button { padding: 8px 12px; border: 1px solid var(--line); border-radius: 5px; background: var(--paper); color: var(--accent); }
    button:focus-visible, input:focus-visible { outline: 2px solid var(--accent); outline-offset: 3px; }
    button:disabled { opacity: .45; cursor: default; }
    .card-actions { justify-content: space-between; padding: 12px 18px; border-top: 1px solid var(--line); }
    .card-actions button[aria-pressed="true"] { background: var(--accent-soft); border-color: var(--accent); }
    .card-back { padding: 24px 18px; min-height: 220px; }
    .card-back .word { margin-bottom: 24px; }
    .card-back .detail-value { line-height: 1.8; }
    .card-head > div:first-child { min-width: 0; }
    .word, #progress-label { overflow-wrap: anywhere; }
    .word-meta { flex-wrap: wrap; }
    .pagination { justify-content: center; padding: 24px 0; }
    #storage-status { color: var(--muted); margin: 10px 0 0; font-size: .75rem; }
    #storage-status.error { color: #a13420; }
    #grid { scroll-margin-top: 340px; }
    .card { scroll-margin-top: 340px; }
    @media (max-width: 640px) { .topbar { position: static; } .stats { white-space: normal; } #grid, .card { scroll-margin-top: 16px; } }

    @media (min-width: 641px) and (max-width: 820px) { .toolbar { grid-template-columns: minmax(0, 1fr) auto; } .stats { grid-column: 1 / -1; justify-content: space-between; } }
    @media (max-width: 640px) { .topbar-inner { padding: 18px 16px 14px; } main { padding: 23px 16px 52px; } .masthead { display: block; } .resume-button { margin-top: 14px; } .toolbar { grid-template-columns: 1fr; gap: 10px; margin-top: 17px; } .voice-control { justify-content: space-between; } .stats { justify-content: space-between; } }
    @media (prefers-reduced-motion: reduce) { *, *::before, *::after { scroll-behavior: auto !important; transition: none !important; } }
  </style>
</head>
<body>
  <header class="topbar">
    <div class="topbar-inner">
      <div class="masthead">
        <div><p class="kicker">Vocabulary studio</p><h1>考研生词</h1><p class="subtitle">按自己的节奏，逐词建立长期记忆</p></div>
        <button id="resume" class="resume-button" type="button" aria-label="回到上次学习位置"><span aria-hidden="true">↗</span> 继续上次学习</button>
      </div>
      <div class="toolbar">
        <div class="search-wrap">
          <svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" aria-hidden="true"><circle cx="11" cy="11" r="7"></circle><path d="m20 20-4-4"></path></svg>
          <input id="search" type="search" autocomplete="off" placeholder="搜索单词、释义或词性…" aria-label="搜索单词、释义或词性">
          <div id="search-results" class="search-results" role="listbox"></div>
        </div>
        <label class="voice-control" for="voice-select"><span>发音</span><select id="voice-select" aria-label="选择发音口音"><option value="en-GB" selected>英式 English (UK)</option><option value="en-US">美式 English (US)</option></select></label>
        <div class="stats"><span><strong id="visible-count">0</strong> / <span id="total-count">0</span> 词</span><span id="progress-label">未开始</span></div>
      </div>
      <div class="study-toolbar">
        <label><input id="review-only" type="checkbox"> 只看需巩固（<span id="review-count">0</span>）</label>
        <button id="export-study" type="button">导出学习记录</button>
        <button id="import-study" type="button">导入学习记录</button>
        <input id="import-file" type="file" accept="application/json,.json" hidden>
      </div>
      <p id="storage-status" role="status">标记自动保存在此浏览器 · 建议定期导出备份</p>
    </div>
  </header>
  <main>
    <div class="section-line"><span id="section-title">全部词汇</span><span id="loaded-count">正在准备…</span></div>
    <section id="grid" aria-live="polite"></section>
    <div id="empty" class="empty-state">没有找到匹配的单词，请换个关键词试试。</div>
    <nav class="pagination" aria-label="词卡分页">
      <button id="previous-page" type="button">上一页</button><span id="page-label"></span><button id="next-page" type="button">下一页</button>
    </nav>
  </main>
  <script>
    const VOCAB = __VOCAB_JSON__;
    const PROGRESS_KEY = 'qwentoword-progress-v1';
    const RECORD_PREFIX = 'vocabulary-study-v2:word:';
    const POSITION_KEY = 'vocabulary-study-v2:position';
    const BATCH_SIZE = 36;
    const wordKey = word => String(word).normalize('NFKC').trim().toLowerCase();
    const byId = new Map(VOCAB.map(item => [item.id, item]));
    const byWord = new Map(VOCAB.map(item => [wordKey(item.word), item]));
    const searchIndex = new Map(VOCAB.map(item => [item.id, [item.id, item.word, item.translate, item.part_of_speech, item.memory_techniques, item.similar_words, item.example_sentence, item.english_definition].join(' ').toLowerCase()]));
    const records = new Map();
    const flipped = new Set();
    const state = { query: '', matches: VOCAB, page: 0, lastId: null, reviewOnly: false };
    const grid = document.getElementById('grid');
    const empty = document.getElementById('empty');
    const search = document.getElementById('search');
    const searchResults = document.getElementById('search-results');
    const voiceSelect = document.getElementById('voice-select');
    const VOICE_KEY = 'qwentoword-voice-v1';
    const speechState = { voices: [] };
    const esc = value => String(value ?? '').replace(/[&<>"']/g, character => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[character]));
    let storageFailed = false;
    function storageNotice(message, error = false) {
      if (error) storageFailed = true;
      const status = document.getElementById('storage-status');
      status.textContent = storageFailed ? '本地保存不完整，请导出备份。' + (error ? message : '') : message;
      status.classList.toggle('error', storageFailed);
    }
    function validRecord(value) {
      return value && typeof value.word === 'string' && value.word.length > 0 && value.word.length <= 500
        && wordKey(value.word).length > 0 && typeof value.review === 'boolean' && Number.isSafeInteger(value.at) && value.at >= 0;
    }
    function validPosition(value) {
      return value && typeof value.word === 'string' && value.word.length <= 500 && Number.isSafeInteger(value.at) && value.at >= 0;
    }
    function readSaved(key, validate) {
      let foundInvalid = false;
      for (const candidate of [key, key + ':backup']) {
        try {
          const raw = localStorage.getItem(candidate);
          if (raw === null) continue;
          const value = JSON.parse(raw);
          if (validate(value)) {
            if (foundInvalid) storageNotice('已从备用记录恢复，请导出备份。');
            return value;
          }
          foundInvalid = true;
        } catch (_) { foundInvalid = true; }
      }
      if (foundInvalid) storageNotice('无法读取部分记录。', true);
      return null;
    }
    function writeSaved(key, value) {
      const data = JSON.stringify(value);
      try {
        localStorage.setItem(key, data);
        localStorage.setItem(key + ':backup', data);
        storageNotice('已保存到此浏览器 · 建议定期导出备份');
        return true;
      } catch (_) {
        storageNotice('当前改动仍在本页，关闭前请导出。', true);
        return false;
      }
    }
    function recordKey(word) { return RECORD_PREFIX + encodeURIComponent(wordKey(word)); }
    function isReview(item) { return records.get(wordKey(item.word))?.review === true; }
    function setReview(item) {
      const key = wordKey(item.word);
      const latest = readSaved(recordKey(key), validRecord);
      const previous = records.get(key);
      const current = latest && (!previous || latest.at > previous.at) ? latest : previous;
      const value = { word: key, review: current?.review !== true, at: Math.max(Date.now(), (previous?.at || 0) + 1, (latest?.at || 0) + 1) };
      records.set(key, value);
      writeSaved(recordKey(key), value);
      saveProgress(item.id);
      setQuery(search.value, false);
      const focusCard = document.getElementById(`word-card-${item.id}`) || grid.firstElementChild;
      (focusCard?.querySelector('[data-review]') || document.getElementById('review-only')).focus({ preventScroll: true });
    }
    function saveProgress(id) {
      const item = byId.get(id);
      if (!item) return;
      state.lastId = id;
      writeSaved(POSITION_KEY, { word: wordKey(item.word), at: Date.now() });
      updateProgress();
    }
    function updateProgress() {
      const item = byId.get(state.lastId);
      document.getElementById('progress-label').textContent = item ? `上次：${item.word}` : '未开始';
      document.getElementById('resume').classList.toggle('visible', Boolean(item));
      for (const card of grid.children) card.classList.toggle('current', card.dataset.id === state.lastId);
    }
    function iconSpeaker() { return '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" aria-hidden="true"><path d="M11 5 6 9H3v6h3l5 4V5Z"></path><path d="M15.5 8.5a5 5 0 0 1 0 7M18.5 5.5a9 9 0 0 1 0 13"></path></svg>'; }
    function refreshVoices() { speechState.voices = "speechSynthesis" in window ? window.speechSynthesis.getVoices() : []; }
    function voiceFor(locale) {
      const normalisedLocale = locale.toLowerCase().replace("_", "-");
      return speechState.voices.find(voice => voice.lang.toLowerCase() === normalisedLocale)
        || speechState.voices.find(voice => voice.lang.toLowerCase().startsWith(`${normalisedLocale}-`));
    }
    function speak(word, rate) {
      if (!("speechSynthesis" in window)) { alert("当前浏览器不支持单词朗读"); return; }
      window.speechSynthesis.cancel();
      const utterance = new SpeechSynthesisUtterance(word);
      const locale = voiceSelect.value || "en-GB";
      utterance.lang = locale; utterance.rate = rate || .86; utterance.pitch = 1;
      const voice = voiceFor(locale);
      if (voice) utterance.voice = voice;
      window.speechSynthesis.speak(utterance);
    }
    function speechText(value) {
      return String(value ?? "").replace(/(?:\s*(?:\([^()（）]*\)|（[^()（）]*）))+\s*([.!?。！？]?)\s*$/, "$1").trim();
    }
    function openDictionary(word) {
      const term = String(word || "").trim();
      if (!term) return;
      window.open(`https://dict.youdao.com/result?word=${encodeURIComponent(term)}&lang=en`, "_blank", "noopener,noreferrer");
    }
    function getPhonetic(item) { return item.phonetic || ""; }
    function cardMarkup(item) {
      const phonetic = getPhonetic(item);
      const phoneticMarkup = phonetic
        ? `<button class="phonetic-button" type="button" data-speak-phonetic="${esc(item.word)}" aria-label="朗读 ${esc(item.word)} 的音标"><span class="phonetic">/${esc(phonetic.replace(/^\/+|\/+$/g, ""))}/</span></button>`
        : `<button class="phonetic-button" type="button" data-speak-phonetic="${esc(item.word)}" aria-label="朗读 ${esc(item.word)}"><span class="phonetic missing">音标待补充</span></button>`;
      const fields = [["英译英", item.english_definition], ["词形变化", item.transformation], ["记忆方法", item.memory_techniques], ["近义辨析", item.similar_words], ["日常例句", item.example_sentence]];
      const detailMarkup = ([label, value]) => {
        const content = value ? (label === "日常例句" || label === "英译英"
          ? `<button class="example-button" type="button" data-speak-example="${esc(value)}" aria-label="朗读${esc(label)}">${iconSpeaker()}${esc(value)}</button>`
          : esc(value)) : "待补充";
        return `<div class="detail"><span class="detail-label">${label}</span><span class="detail-value${value ? "" : " empty"}">${content}</span></div>`;
      };
      return `<article class="card" id="word-card-${esc(item.id)}" data-id="${esc(item.id)}"><div class="card-front"${flipped.has(item.id) ? " hidden" : ""}><div class="card-head"><div><h2 class="word" data-word="${esc(item.word)}" tabindex="0" title="双击查询有道词典">${esc(item.word)}</h2><div class="word-meta">${phoneticMarkup}<span aria-hidden="true">·</span><span>${esc(item.part_of_speech) || "词性待补充"}</span></div></div><div class="badges"><span class="word-id">${esc(item.id)}</span><button class="speak" type="button" data-speak="${esc(item.word)}" aria-label="朗读 ${esc(item.word)}" title="朗读单词">${iconSpeaker()}</button>${item.number ? `<div class="number">${esc(item.number)}</div>` : ""}</div></div>${item.translate ? `<p class="meaning">${esc(item.translate)}</p>` : `<p class="meaning detail-value empty">释义待补充</p>`}<div class="details">${fields.map(detailMarkup).join("")}</div></div><div class="card-back"${flipped.has(item.id) ? "" : " hidden"}><h2 class="word" data-word="${esc(item.word)}">${esc(item.word)}</h2><div class="detail-value">${item.example_sentence ? `<button class="example-button" type="button" data-speak-example="${esc(item.example_sentence)}" aria-label="朗读日常例句">${esc(item.example_sentence)}</button>` : "例句待补充"}</div></div><div class="card-actions"><button type="button" data-review="${esc(item.id)}" aria-pressed="${isReview(item)}">${isReview(item) ? "✓ 需巩固" : "标记需巩固"}</button><button type="button" data-flip="${esc(item.id)}" aria-pressed="${flipped.has(item.id)}">${flipped.has(item.id) ? "查看释义" : "翻到背面"}</button></div></article>`;
    }
    function renderPage() {
      state.page = Math.max(0, Math.min(state.page, Math.ceil(state.matches.length / BATCH_SIZE) - 1));
      const start = state.page * BATCH_SIZE;
      grid.innerHTML = state.matches.slice(start, start + BATCH_SIZE).map(cardMarkup).join('');
      document.getElementById('visible-count').textContent = state.matches.length;
      document.getElementById('total-count').textContent = VOCAB.length;
      document.getElementById('review-count').textContent = VOCAB.filter(isReview).length;
      document.getElementById('loaded-count').textContent = state.matches.length ? `${start + 1}–${Math.min(start + BATCH_SIZE, state.matches.length)} / ${state.matches.length}` : '0 / 0';
      document.getElementById('page-label').textContent = `${state.page + 1} / ${Math.max(1, Math.ceil(state.matches.length / BATCH_SIZE))} 页`;
      document.getElementById('previous-page').disabled = state.page === 0;
      document.getElementById('next-page').disabled = start + BATCH_SIZE >= state.matches.length;
      document.getElementById('section-title').textContent = (state.reviewOnly ? '需巩固词汇' : '全部词汇') + (state.query ? ` · 搜索“${state.query}”` : '');
      empty.classList.toggle('visible', !state.matches.length);
      empty.textContent = state.reviewOnly ? '没有符合条件的需巩固单词。可关闭筛选，在词卡底部标记。' : '没有找到匹配的单词，请换个关键词试试。';
      updateProgress();
    }
    function matchingItems(query) {
      const normalized = query.trim().toLowerCase();
      return VOCAB.filter(item => (!state.reviewOnly || isReview(item)) && (!normalized || searchIndex.get(item.id).includes(normalized)));
    }
    function setQuery(query, resetPage = true) {
      state.query = query.trim();
      state.matches = matchingItems(query);
      if (resetPage) state.page = 0;
      renderPage();
      searchResults.classList.remove('open');
    }
    function searchItems(query) {
      const normalized = query.trim().toLowerCase();
      if (!normalized) return [];
      const matches = matchingItems(query);
      return matches.filter(item => item.word.toLowerCase() === normalized).concat(matches.filter(item => item.word.toLowerCase() !== normalized));
    }
    function showResults(query) {
      const term = query.trim();
      const matches = searchItems(term).slice(0, 8);
      searchResults.innerHTML = matches.length
        ? matches.map(item => `<button class="search-result" type="button" data-result-id="${esc(item.id)}"><span>${esc(item.word)}</span><small>${esc(item.translate || '')}</small></button>`).join('')
        : term ? `<button class="search-result" type="button" data-dictionary-word="${esc(term)}">在有道词典查询 ${esc(term)}</button>` : '';
      searchResults.classList.toggle('open', Boolean(term));
    }
    function locate(id) {
      const index = state.matches.findIndex(item => item.id === id);
      if (index < 0) return;
      state.page = Math.floor(index / BATCH_SIZE);
      renderPage();
      const target = document.getElementById(`word-card-${id}`);
      target?.scrollIntoView({ block: 'start' });
      saveProgress(id);
    }
    function turnPage(direction) {
      state.page += direction;
      renderPage();
      const first = state.matches[state.page * BATCH_SIZE];
      if (first) saveProgress(first.id);
      grid.scrollIntoView({ block: 'start' });
    }
    // Shared by both card faces: pronunciation and dictionary controls never flip a card.
    document.addEventListener('click', event => {
      const card = event.target.closest('.card');
      if (card) saveProgress(card.dataset.id);
      const review = event.target.closest('[data-review]');
      if (review) { setReview(byId.get(review.dataset.review)); return; }
      const flip = event.target.closest('[data-flip]');
      if (flip) {
        const id = flip.dataset.flip;
        if (flipped.has(id)) flipped.delete(id); else flipped.add(id);
        card.querySelector('.card-front').hidden = flipped.has(id);
        card.querySelector('.card-back').hidden = !flipped.has(id);
        flip.setAttribute('aria-pressed', String(flipped.has(id)));
        flip.textContent = flipped.has(id) ? '查看释义' : '翻到背面';
        return;
      }
      const speakButton = event.target.closest('[data-speak]');
      if (speakButton) speak(speakButton.dataset.speak);
      const exampleButton = event.target.closest('[data-speak-example]');
      if (exampleButton) speak(speechText(exampleButton.dataset.speakExample), .8);
      const ipaButton = event.target.closest('[data-speak-phonetic]');
      if (ipaButton) speak(ipaButton.dataset.speakPhonetic, .63);
      const result = event.target.closest('[data-result-id]');
      if (result) { searchResults.classList.remove('open'); locate(result.dataset.resultId); }
      const dictionary = event.target.closest('[data-dictionary-word]');
      if (dictionary) { openDictionary(dictionary.dataset.dictionaryWord); searchResults.classList.remove('open'); }
      if (!event.target.closest('.search-wrap')) searchResults.classList.remove('open');
    });
    document.addEventListener('dblclick', event => { const word = event.target.closest('[data-word]'); if (word) openDictionary(word.dataset.word); });
    search.addEventListener('input', () => { setQuery(search.value); showResults(search.value); });
    search.addEventListener('keydown', event => {
      if (event.isComposing) return;
      if (event.key === 'Escape') { searchResults.classList.remove('open'); search.blur(); }
      if (event.key === 'Enter' && search.value.trim()) {
        event.preventDefault();
        searchResults.classList.remove('open');
        const first = searchItems(search.value)[0];
        if (first) locate(first.id); else openDictionary(search.value);
      }
    });
    document.getElementById('review-only').addEventListener('change', event => { state.reviewOnly = event.target.checked; setQuery(search.value); });
    document.getElementById('previous-page').addEventListener('click', () => turnPage(-1));
    document.getElementById('next-page').addEventListener('click', () => turnPage(1));
    document.getElementById('resume').addEventListener('click', () => {
      const id = state.lastId;
      state.reviewOnly = false;
      document.getElementById('review-only').checked = false;
      search.value = '';
      setQuery('');
      if (id) locate(id);
    });
    function backupData() {
      const item = byId.get(state.lastId);
      return { format: 'vocabulary-study', version: 2, exportedAt: new Date().toISOString(), records: [...records.values()], position: item ? { word: wordKey(item.word), at: Date.now() } : null };
    }
    function importData(data) {
      if (!data || data.format !== 'vocabulary-study' || data.version !== 2 || !Array.isArray(data.records) || data.records.length > 100000
        || !data.records.every(validRecord) || (data.position !== null && !validPosition(data.position))) throw new Error('备份格式无效或版本不支持');
      // Validate the whole file before touching existing records. Newer marks win, including removals.
      for (const entry of data.records) {
        const key = wordKey(entry.word);
        if (!key) throw new Error('备份包含空白单词');
      }
      for (const entry of data.records) {
        const key = wordKey(entry.word);
        const current = records.get(key);
        const disk = readSaved(recordKey(key), validRecord);
        const latest = disk && (!current || disk.at > current.at) ? disk : current;
        if (!latest || entry.at > latest.at) {
          const record = { word: key, review: entry.review, at: entry.at };
          records.set(key, record);
          writeSaved(recordKey(key), record);
        } else if (latest) records.set(key, latest);
      }
      const saved = readSaved(POSITION_KEY, validPosition);
      if (data.position && (!saved || data.position.at > saved.at)) {
        writeSaved(POSITION_KEY, data.position);
        state.lastId = byWord.get(wordKey(data.position.word))?.id || state.lastId;
      }
      setQuery(search.value, false);
    }
    document.getElementById('export-study').addEventListener('click', () => {
      const url = URL.createObjectURL(new Blob([JSON.stringify(backupData(), null, 2)], { type: 'application/json' }));
      const link = document.createElement('a');
      link.href = url; link.download = `词汇学习记录-${new Date().toISOString().slice(0, 10)}.json`;
      document.body.appendChild(link); link.click(); link.remove();
      setTimeout(() => URL.revokeObjectURL(url), 1000);
    });
    const importInput = document.getElementById('import-file');
    document.getElementById('import-study').addEventListener('click', () => importInput.click());
    importInput.addEventListener('change', async () => {
      const file = importInput.files[0];
      if (!file) return;
      try {
        if (file.size > 20 * 1024 * 1024) throw new Error('备份文件超过 20 MB');
        importData(JSON.parse(await file.text()));
        storageNotice('备份已合并，已有的较新标记会保留。');
      } catch (error) { storageNotice(`导入失败：${error.message}`); }
      finally { importInput.value = ''; }
    });
    window.addEventListener('storage', event => {
      if (event.key?.startsWith(RECORD_PREFIX) && !event.key.endsWith(':backup') && event.newValue) {
        try {
          const record = JSON.parse(event.newValue);
          const current = records.get(wordKey(record.word));
          if (validRecord(record) && (!current || record.at >= current.at)) { records.set(wordKey(record.word), record); setQuery(search.value, false); }
        } catch (_) { storageNotice('其他窗口的记录无法读取。', true); }
      }
      if (event.key === null || (event.key?.startsWith(RECORD_PREFIX) && !event.newValue)) storageNotice('浏览器中的记录已被清除，请导出当前记录备份。', true);
    });
    function initializeStudy() {
      try {
        const keys = new Set();
        for (let i = 0; i < localStorage.length; i++) {
          const key = localStorage.key(i);
          if (key?.startsWith(RECORD_PREFIX)) keys.add(key.replace(/:backup$/, ''));
        }
        for (const key of keys) { const value = readSaved(key, validRecord); if (value) records.set(wordKey(value.word), value); }
        const position = readSaved(POSITION_KEY, validPosition);
        if (position) state.lastId = byWord.get(wordKey(position.word))?.id || null;
        else {
          const legacy = JSON.parse(localStorage.getItem(PROGRESS_KEY) || 'null');
          if (legacy && byId.has(legacy.id)) { state.lastId = legacy.id; writeSaved(POSITION_KEY, { word: wordKey(byId.get(legacy.id).word), at: Date.now() }); }
        }
        const voice = localStorage.getItem(VOICE_KEY);
        if (['en-GB', 'en-US'].includes(voice)) voiceSelect.value = voice;
      } catch (_) { storageNotice('无法读取本地记录。', true); }
      voiceSelect.addEventListener('change', () => { try { localStorage.setItem(VOICE_KEY, voiceSelect.value); } catch (_) { storageNotice('口音设置未保存。', true); } });
      refreshVoices();
      if ('speechSynthesis' in window) window.speechSynthesis.addEventListener('voiceschanged', refreshVoices);
      if (state.lastId) state.page = Math.floor(VOCAB.findIndex(item => item.id === state.lastId) / BATCH_SIZE);
      renderPage();
    }
    initializeStudy();
  </script>
</body>
</html>
'''


def generate(
    input_path: Path,
    output_path: Path,
    sheet_name: str | None = None,
    words_path: Path | None = None,
) -> int:
    phonetic_path = words_path or input_path.with_name("words.json")
    update_workbook(input_path, phonetic_path, sheet_name)
    rows = load_rows(input_path, sheet_name)
    serialised = (
        json.dumps(rows, ensure_ascii=False, separators=(",", ":"))
        .replace("<", "\\u003c")
        .replace("\u2028", "\\u2028")
        .replace("\u2029", "\\u2029")
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(HTML_TEMPLATE.replace("__VOCAB_JSON__", serialised), encoding="utf-8")
    return len(rows)


def main() -> None:
    args = parse_args()
    output = args.output or args.input.with_suffix(".html")
    words_path = args.words or args.input.with_name("words.json")
    count = generate(args.input, output, args.sheet, words_path)
    print(f"Generated {output} with {count} vocabulary cards.")


if __name__ == "__main__":
    main()

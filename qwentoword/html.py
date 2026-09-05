"""Update vocabulary metadata and generate a browser-based reader.

Run ``python html.py`` to add unique IDs and local ``words.json`` phonetics to
``考研生词.xlsx``, then create ``考研生词.html``.  The generated document is
self-contained: workbook data is embedded as JSON and all interaction is
implemented in the page itself, so it can be opened directly from disk.
"""

from __future__ import annotations

import argparse
import json
import os
import re
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree


FIELD_ALIASES = {
    "id": ("id", "word id", "word_id", "单词编号", "唯一编号"),
    "word": ("word", "单词", "词汇"),
    "translate": ("translate", "translation", "释义", "翻译", "中文释义"),
    "number": ("number", "编号", "频次", "frequency"),
    "part_of_speech": ("part of speech", "part_of_speech", "词性"),
    "transformation": ("transformation", "word forms", "词形变化"),
    "memory_techniques": ("memory techniques", "memory_technique", "记忆方法"),
    "similar_words": (
        "distinguishing between similar words",
        "similar words",
        "近义词辨析",
        "易混淆词",
    ),
    "phonetic": (
        "phonetic",
        "phonetics",
        "ipa",
        "音标",
        "英式音标",
        "美式音标",
    ),
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Add vocabulary IDs/phonetics to Excel and generate its study webpage.")
    parser.add_argument("--input", type=Path, default=Path(os.getenv("VOCAB_INPUT", "考研生词.xlsx")))
    parser.add_argument("--output", type=Path, default=None)
    parser.add_argument("--sheet", default=None, help="Worksheet name; defaults to the active worksheet.")
    parser.add_argument("--words", type=Path, default=None, help="Phonetic JSON path; defaults to words.json beside the workbook.")
    return parser.parse_args()


def normalise(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).lower()


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def find_columns(headers: list[Any]) -> dict[str, int | None]:
    by_name = {normalise(value): i for i, value in enumerate(headers) if value is not None}
    result: dict[str, int | None] = {}
    for field, aliases in FIELD_ALIASES.items():
        result[field] = next((by_name[normalise(alias)] for alias in aliases if normalise(alias) in by_name), None)
    if result["word"] is None:
        raise ValueError("Workbook must contain a 'word' (or 单词) column")
    return result


def load_phonetics(path: Path) -> dict[str, str]:
    """Load ``word -> phonetic`` entries from the local words.json file."""

    if not path.exists():
        return {}
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"Could not read phonetic JSON: {path}") from exc

    entries: Any
    if isinstance(payload, list):
        entries = payload
    elif isinstance(payload, dict):
        # Also accept a compact {"word": "/.../"} mapping.
        entries = [
            {"word": word, "phonetic": phonetic}
            for word, phonetic in payload.items()
            if isinstance(word, str)
        ]
    else:
        raise ValueError("Phonetic JSON must be a list or an object mapping words to phonetics")

    result: dict[str, str] = {}
    for entry in entries:
        if not isinstance(entry, dict):
            continue
        word = cell_text(entry.get("word"))
        phonetic = cell_text(entry.get("phonetic"))
        if word and phonetic:
            result.setdefault(word.casefold(), phonetic)
    return result


def next_word_id(used_ids: set[str]) -> str:
    """Return a readable, stable-looking ID that cannot collide in the sheet."""

    numbers = []
    for value in used_ids:
        match = re.fullmatch(r"W(\d+)", value, flags=re.IGNORECASE)
        if match:
            numbers.append(int(match.group(1)))
    candidate = max(numbers, default=0) + 1
    while f"W{candidate:06d}" in used_ids:
        candidate += 1
    return f"W{candidate:06d}"


def _rows_from_values(headers: list[Any], data_rows: list[tuple[int, list[Any]]]) -> list[dict[str, str]]:
    """Convert worksheet values into the small, stable JSON schema used by the page."""

    columns = find_columns(headers)
    fields = ("translate", "number", "part_of_speech", "transformation", "memory_techniques", "similar_words", "phonetic")
    rows: list[dict[str, str]] = []
    for source_index, values in data_rows:
        word_column = columns["word"]
        word = cell_text(values[word_column]) if word_column is not None and word_column < len(values) else ""
        if not word:
            continue
        id_column = columns["id"]
        row_id = cell_text(values[id_column]) if id_column is not None and id_column < len(values) else ""
        row = {"id": row_id or f"W{len(rows) + 1:06d}", "sourceRow": str(source_index), "word": word}
        for field in fields:
            column = columns[field]
            row[field] = cell_text(values[column]) if column is not None and column < len(values) else ""
        rows.append(row)
    return rows


def _load_rows_with_xml(path: Path, sheet_name: str | None = None) -> list[dict[str, str]]:
    """Read normal XLSX cell types without third-party packages.

    This deliberately covers the worksheet features needed for a vocabulary
    table (inline strings, shared strings, numbers, booleans and blanks).  If
    openpyxl is available it remains the preferred parser.
    """

    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ns = {"m": main_ns, "r": rel_ns, "pr": package_rel_ns}
    with ZipFile(path) as archive:
        workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        sheets = workbook.find("m:sheets", ns)
        if sheets is None:
            raise ValueError("Workbook does not contain any worksheets")
        sheet_nodes = list(sheets)
        selected = None
        if sheet_name:
            selected = next((node for node in sheet_nodes if node.attrib.get("name") == sheet_name), None)
            if selected is None:
                raise ValueError(f"Worksheet not found: {sheet_name}")
        else:
            active_tab = 0
            views = workbook.find("m:bookViews", ns)
            if views is not None:
                view = views.find("m:workbookView", ns)
                if view is not None:
                    try:
                        active_tab = max(0, int(view.attrib.get("activeTab", "0")))
                    except ValueError:
                        active_tab = 0
            selected = sheet_nodes[min(active_tab, len(sheet_nodes) - 1)]

        relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        relationship_map = {node.attrib.get("Id"): node.attrib.get("Target", "") for node in relationships}
        target = relationship_map.get(selected.attrib.get(f"{{{rel_ns}}}id"), "")
        target = target.lstrip("/")
        if not target.startswith("xl/"):
            target = f"xl/{target}"
        sheet_root = ElementTree.fromstring(archive.read(target))

        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in shared_root.findall("m:si", ns):
                shared_strings.append("".join(item.itertext()))

        def read_cell(cell: ElementTree.Element) -> str:
            value = cell.find("m:v", ns)
            cell_type = cell.attrib.get("t", "")
            if cell_type == "inlineStr":
                inline = cell.find("m:is", ns)
                return "".join(inline.itertext()) if inline is not None else ""
            if value is None or value.text is None:
                return ""
            raw = value.text
            if cell_type == "s":
                try:
                    return shared_strings[int(raw)]
                except (ValueError, IndexError):
                    return ""
            if cell_type == "b":
                return "TRUE" if raw == "1" else "FALSE"
            return raw

        def column_index(reference: str) -> int:
            letters = re.match(r"[A-Za-z]+", reference or "")
            if not letters:
                return 0
            index = 0
            for letter in letters.group(0).upper():
                index = index * 26 + ord(letter) - ord("A") + 1
            return index - 1

        sheet_data = sheet_root.find("m:sheetData", ns)
        if sheet_data is None:
            raise ValueError("Worksheet is empty")
        parsed_rows: list[tuple[int, dict[int, str]]] = []
        for row_node in sheet_data.findall("m:row", ns):
            row_number = int(row_node.attrib.get("r", len(parsed_rows) + 1))
            cells = {column_index(cell.attrib.get("r", "")): read_cell(cell) for cell in row_node.findall("m:c", ns)}
            parsed_rows.append((row_number, cells))
        if not parsed_rows:
            raise ValueError("Worksheet is empty")
        header_map = parsed_rows[0][1]
        max_header_column = max(header_map, default=-1)
        headers = [header_map.get(index, "") for index in range(max_header_column + 1)]
        data_rows = []
        for row_number, values in parsed_rows[1:]:
            width = max(max_header_column + 1, max(values, default=-1) + 1)
            data_rows.append((row_number, [values.get(index, "") for index in range(width)]))
    return _rows_from_values(headers, data_rows)


def load_rows(path: Path, sheet_name: str | None = None) -> list[dict[str, str]]:
    try:
        import openpyxl
    except ImportError:
        return _load_rows_with_xml(path, sheet_name)

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = list(next(iterator))
        except StopIteration as exc:
            raise ValueError("Workbook is empty") from exc
        return _rows_from_values(headers, list(enumerate(iterator, start=2)))
    finally:
        workbook.close()


def _update_workbook_with_openpyxl(path: Path, phonetics: dict[str, str], sheet_name: str | None) -> None:
    import openpyxl

    workbook = openpyxl.load_workbook(path)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        headers = [cell.value for cell in sheet[1]]
        columns = find_columns(headers)
        word_column = columns["word"]
        assert word_column is not None
        next_column = sheet.max_column
        id_column = columns["id"]
        if id_column is None:
            id_column = next_column
            next_column += 1
        phonetic_column = columns["phonetic"]
        if phonetic_column is None:
            phonetic_column = next_column
        if columns["id"] is None:
            sheet.cell(1, id_column + 1).value = "id"
        if columns["phonetic"] is None:
            sheet.cell(1, phonetic_column + 1).value = "phonetic"

        used_ids: set[str] = set()
        for row_number in range(2, sheet.max_row + 1):
            word = cell_text(sheet.cell(row_number, word_column + 1).value)
            if not word:
                continue
            id_cell = sheet.cell(row_number, id_column + 1)
            current_id = cell_text(id_cell.value)
            if not current_id or current_id in used_ids:
                current_id = next_word_id(used_ids)
                id_cell.value = current_id
            used_ids.add(current_id)

            phonetic_cell = sheet.cell(row_number, phonetic_column + 1)
            if not cell_text(phonetic_cell.value):
                phonetic = phonetics.get(word.casefold(), "")
                if phonetic:
                    phonetic_cell.value = phonetic
        workbook.save(path)
    finally:
        workbook.close()


def _xml_column_letters(index: int) -> str:
    result = ""
    value = index + 1
    while value:
        value, remainder = divmod(value - 1, 26)
        result = chr(ord("A") + remainder) + result
    return result


def _update_workbook_with_xml(path: Path, phonetics: dict[str, str], sheet_name: str | None) -> None:
    """Update metadata in an XLSX when openpyxl is unavailable."""

    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ns = {"m": main_ns, "r": rel_ns, "pr": package_rel_ns}
    temp_path = path.with_name(path.name + ".tmp")
    try:
        with ZipFile(path) as archive:
            workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            sheets = workbook.find("m:sheets", ns)
            if sheets is None or not list(sheets):
                raise ValueError("Workbook does not contain any worksheets")
            sheet_nodes = list(sheets)
            if sheet_name:
                selected = next((node for node in sheet_nodes if node.attrib.get("name") == sheet_name), None)
                if selected is None:
                    raise ValueError(f"Worksheet not found: {sheet_name}")
            else:
                active_tab = 0
                views = workbook.find("m:bookViews", ns)
                if views is not None:
                    view = views.find("m:workbookView", ns)
                    if view is not None:
                        try:
                            active_tab = max(0, int(view.attrib.get("activeTab", "0")))
                        except ValueError:
                            active_tab = 0
                selected = sheet_nodes[min(active_tab, len(sheet_nodes) - 1)]

            relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            relationship_map = {node.attrib.get("Id"): node.attrib.get("Target", "") for node in relationships}
            target = relationship_map.get(selected.attrib.get(f"{{{rel_ns}}}id"), "").lstrip("/")
            if not target.startswith("xl/"):
                target = f"xl/{target}"
            sheet_root = ElementTree.fromstring(archive.read(target))
            sheet_data = sheet_root.find("m:sheetData", ns)
            if sheet_data is None:
                raise ValueError("Worksheet is empty")
            sheet_rows = list(sheet_data.findall("m:row", ns))
            if not sheet_rows:
                raise ValueError("Worksheet is empty")

            shared_strings: list[str] = []
            if "xl/sharedStrings.xml" in archive.namelist():
                shared_root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
                shared_strings = ["".join(item.itertext()) for item in shared_root.findall("m:si", ns)]

            def column_index(reference: str) -> int:
                letters = re.match(r"[A-Za-z]+", reference or "")
                if not letters:
                    return 0
                index = 0
                for letter in letters.group(0).upper():
                    index = index * 26 + ord(letter) - ord("A") + 1
                return index - 1

            def read_cell(cell: ElementTree.Element) -> str:
                value = cell.find("m:v", ns)
                cell_type = cell.attrib.get("t", "")
                if cell_type == "inlineStr":
                    inline = cell.find("m:is", ns)
                    return "".join(inline.itertext()) if inline is not None else ""
                if value is None or value.text is None:
                    return ""
                if cell_type == "s":
                    try:
                        return shared_strings[int(value.text)]
                    except (ValueError, IndexError):
                        return ""
                return value.text

            def row_number(row: ElementTree.Element, fallback: int) -> int:
                try:
                    return int(row.attrib.get("r", fallback))
                except ValueError:
                    return fallback

            header_row = sheet_rows[0]
            header_cells = {column_index(cell.attrib.get("r", "")): read_cell(cell) for cell in header_row.findall("m:c", ns)}
            max_column = max((column_index(cell.attrib.get("r", "")) for row in sheet_rows for cell in row.findall("m:c", ns)), default=-1)
            headers = [header_cells.get(index, "") for index in range(max(max_column, max(header_cells, default=-1)) + 1)]
            columns = find_columns(headers)
            word_column = columns["word"]
            assert word_column is not None
            id_column = columns["id"]
            if id_column is None:
                max_column += 1
                id_column = max_column
            phonetic_column = columns["phonetic"]
            if phonetic_column is None:
                max_column += 1
                phonetic_column = max_column

            def find_or_create_cell(row: ElementTree.Element, column: int, number: int) -> ElementTree.Element:
                for cell in row.findall("m:c", ns):
                    if column_index(cell.attrib.get("r", "")) == column:
                        return cell
                cell = ElementTree.Element(f"{{{main_ns}}}c", {"r": f"{_xml_column_letters(column)}{number}"})
                children = row.findall("m:c", ns)
                insert_at = len(children)
                for index, existing in enumerate(children):
                    if column_index(existing.attrib.get("r", "")) > column:
                        insert_at = index
                        break
                row.insert(insert_at, cell)
                return cell

            def set_inline_string(cell: ElementTree.Element, value: str) -> None:
                cell.attrib["t"] = "inlineStr"
                for child in list(cell):
                    cell.remove(child)
                inline = ElementTree.SubElement(cell, f"{{{main_ns}}}is")
                text = ElementTree.SubElement(inline, f"{{{main_ns}}}t")
                text.text = value

            first_number = row_number(header_row, 1)
            set_inline_string(find_or_create_cell(header_row, id_column, first_number), "id")
            set_inline_string(find_or_create_cell(header_row, phonetic_column, first_number), "phonetic")
            used_ids: set[str] = set()
            for position, row in enumerate(sheet_rows[1:], start=2):
                number = row_number(row, position)
                cells = {column_index(cell.attrib.get("r", "")): cell for cell in row.findall("m:c", ns)}
                word = read_cell(cells[word_column]) if word_column in cells else ""
                if not word:
                    continue
                id_cell = find_or_create_cell(row, id_column, number)
                current_id = read_cell(id_cell)
                if not current_id or current_id in used_ids:
                    current_id = next_word_id(used_ids)
                    set_inline_string(id_cell, current_id)
                used_ids.add(current_id)
                phonetic_cell = find_or_create_cell(row, phonetic_column, number)
                if not read_cell(phonetic_cell):
                    phonetic = phonetics.get(word.casefold(), "")
                    if phonetic:
                        set_inline_string(phonetic_cell, phonetic)

            dimension = sheet_root.find("m:dimension", ns)
            if dimension is not None:
                last_row = max(row_number(row, index + 1) for index, row in enumerate(sheet_rows))
                dimension.attrib["ref"] = f"A1:{_xml_column_letters(max_column)}{last_row}"
            ElementTree.register_namespace("", main_ns)
            ElementTree.register_namespace("r", rel_ns)
            updated_sheet = ElementTree.tostring(sheet_root, encoding="utf-8", xml_declaration=True)
            with ZipFile(temp_path, "w", ZIP_DEFLATED) as output:
                for item in archive.infolist():
                    output.writestr(item, updated_sheet if item.filename == target else archive.read(item.filename))
        os.replace(temp_path, path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def update_workbook(path: Path, words_path: Path, sheet_name: str | None = None) -> None:
    phonetics = load_phonetics(words_path)
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        _update_workbook_with_xml(path, phonetics, sheet_name)
    else:
        _update_workbook_with_openpyxl(path, phonetics, sheet_name)


HTML_TEMPLATE = r'''<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <meta name="theme-color" content="#18252b">
  <title>考研生词 · 学习卡片</title>
  <style>
    :root { color-scheme: light; --ink: #18252b; --muted: #68757a; --line: #dce5e5; --paper: #fff; --wash: #f4f7f6; --accent: #0c766d; --accent-soft: #e5f3f0; --gold: #c9872c; }
    * { box-sizing: border-box; }
    html { scroll-behavior: smooth; }
    body { margin: 0; min-width: 320px; background: var(--wash); color: var(--ink); font-family: "Noto Sans SC", "PingFang SC", "Microsoft YaHei", system-ui, sans-serif; }
    button, input { font: inherit; }
    button { cursor: pointer; }
    .topbar { position: sticky; z-index: 10; top: 0; background: rgba(244,247,246,.94); border-bottom: 1px solid var(--line); backdrop-filter: blur(16px); }
    .topbar-inner { max-width: 1360px; margin: 0 auto; padding: 22px 30px 18px; }
    .masthead { display: flex; align-items: flex-start; justify-content: space-between; gap: 20px; }
    h1 { margin: 0; font-size: clamp(1.55rem, 3vw, 2.45rem); letter-spacing: .01em; line-height: 1.1; }
    .kicker { margin: 0 0 8px; color: var(--accent); font-size: .74rem; font-weight: 800; letter-spacing: .13em; text-transform: uppercase; }
    .subtitle { margin: 9px 0 0; color: var(--muted); font-size: .9rem; }
    .resume-button { display: none; border: 1px solid var(--accent); border-radius: 5px; padding: 9px 13px; color: var(--accent); background: var(--paper); font-weight: 700; white-space: nowrap; }
    .resume-button.visible { display: inline-flex; align-items: center; gap: 7px; }
    .resume-button:hover { background: var(--accent-soft); }
    .toolbar { display: grid; grid-template-columns: minmax(200px, 1fr) auto auto; align-items: center; gap: 15px; margin-top: 22px; }
    .search-wrap { position: relative; }
    .search-wrap svg { position: absolute; left: 15px; top: 50%; width: 18px; height: 18px; color: var(--muted); transform: translateY(-50%); pointer-events: none; }
    #search { width: 100%; border: 1px solid #cbd8d7; border-radius: 5px; padding: 13px 15px 13px 44px; outline: 0; background: var(--paper); color: var(--ink); }
    #search:focus { border-color: var(--accent); box-shadow: 0 0 0 3px rgba(12,118,109,.12); }
    .voice-control { display: inline-flex; align-items: center; gap: 8px; border: 1px solid #cbd8d7; border-radius: 5px; padding: 8px 10px; background: var(--paper); color: var(--muted); font-size: .78rem; white-space: nowrap; }
    #voice-select { border: 0; outline: 0; background: transparent; color: var(--ink); font-weight: 700; cursor: pointer; }
    #voice-select:focus-visible { box-shadow: 0 0 0 3px rgba(12,118,109,.12); }
    .stats { display: flex; align-items: center; justify-content: flex-end; gap: 16px; color: var(--muted); font-size: .82rem; white-space: nowrap; }
    .stats strong { color: var(--ink); font-size: 1rem; }
    .meter { width: 110px; height: 6px; overflow: hidden; border-radius: 9px; background: #d8e4e2; }
    .meter i { display: block; width: 0; height: 100%; border-radius: inherit; background: var(--accent); transition: width .25s ease; }
    .search-results { position: absolute; z-index: 11; top: calc(100% + 6px); left: 0; right: 0; display: none; overflow: hidden; border: 1px solid var(--line); border-radius: 5px; background: var(--paper); box-shadow: 0 12px 30px rgba(24,37,43,.12); }
    .search-results.open { display: block; }
    .search-result { display: flex; width: 100%; align-items: baseline; justify-content: space-between; gap: 12px; border: 0; border-bottom: 1px solid #edf2f1; padding: 11px 14px; text-align: left; background: transparent; color: var(--ink); }
    .search-result:last-child { border-bottom: 0; }
    .search-result:hover, .search-result:focus-visible { background: var(--accent-soft); outline: 0; }
    .search-result small { color: var(--muted); }
    main { max-width: 1360px; margin: 0 auto; padding: 28px 30px 70px; }
    .section-line { display: flex; align-items: center; justify-content: space-between; margin: 0 0 17px; color: var(--muted); font-size: .78rem; }
    .section-line span:first-child { font-weight: 700; letter-spacing: .08em; text-transform: uppercase; }
    /* Grid auto-flow is intentionally row-major: frequency order reads left-to-right. */
    #grid { display: grid; grid-template-columns: minmax(0, 1fr); align-items: start; gap: 18px; }
    @media (min-width: 650px) { #grid { grid-template-columns: repeat(2, minmax(0, 1fr)); } }
    @media (min-width: 980px) { #grid { grid-template-columns: repeat(3, minmax(0, 1fr)); } }
    @media (min-width: 1320px) { #grid { grid-template-columns: repeat(4, minmax(0, 1fr)); } }
    .card { display: block; width: 100%; margin: 0; overflow: hidden; border: 1px solid var(--line); border-radius: 7px; background: var(--paper); box-shadow: 0 2px 4px rgba(24,37,43,.025); transition: border-color .2s, box-shadow .2s, transform .2s; scroll-margin-top: 148px; }
    .card:hover { border-color: #b8cfcb; box-shadow: 0 8px 22px rgba(24,37,43,.08); transform: translateY(-1px); }
    .card.current { border-color: var(--accent); box-shadow: 0 0 0 3px rgba(12,118,109,.10); }
    .card-head { display: flex; align-items: flex-start; justify-content: space-between; gap: 12px; padding: 18px 18px 13px; }
    .word { margin: 0; font-family: Georgia, "Times New Roman", serif; font-size: 1.75rem; letter-spacing: .01em; line-height: 1; cursor: help; }
    .word:focus-visible { outline: 2px solid var(--accent); outline-offset: 4px; border-radius: 2px; }
    .word-meta { display: flex; align-items: center; gap: 7px; margin-top: 8px; color: var(--muted); font-size: .82rem; }
    .phonetic { color: var(--accent); font-family: Georgia, serif; font-size: .95rem; }
    .phonetic.missing { color: #9ca9a7; font-family: inherit; font-size: .76rem; }
    .speak { display: inline-grid; flex: 0 0 auto; place-items: center; width: 34px; height: 34px; border: 1px solid #cbd8d7; border-radius: 50%; background: transparent; color: var(--accent); }
    .speak:hover, .speak:focus-visible { border-color: var(--accent); background: var(--accent-soft); outline: 0; }
    .speak svg { width: 16px; height: 16px; }
    .phonetic-button { display: inline-flex; align-items: center; gap: 4px; border: 0; padding: 0; color: inherit; background: transparent; }
    .phonetic-button:hover .phonetic { text-decoration: underline; text-underline-offset: 3px; }
    .badges { display: flex; flex-direction: column; align-items: flex-end; gap: 7px; }
    .word-id { color: var(--muted); font-family: ui-monospace, SFMono-Regular, Menlo, monospace; font-size: .68rem; letter-spacing: .04em; }
    .number { flex: 0 0 auto; border-radius: 4px; padding: 4px 7px; color: var(--gold); background: #fff7e9; font-size: .72rem; font-weight: 800; }
    .meaning { margin: 0 18px 15px; border-left: 3px solid var(--accent); padding: 3px 0 3px 11px; color: #263b3e; font-size: .93rem; line-height: 1.65; }
    .details { border-top: 1px solid #edf2f1; padding: 3px 18px 16px; }
    .detail { display: grid; grid-template-columns: 93px minmax(0, 1fr); gap: 9px; padding-top: 12px; font-size: .82rem; line-height: 1.65; }
    .detail-label { color: var(--muted); font-size: .73rem; font-weight: 800; letter-spacing: .05em; }
    .detail-value { min-width: 0; white-space: pre-wrap; overflow-wrap: anywhere; color: #405154; }
    .detail-value.empty { color: #a4afad; font-style: italic; }
    .empty-state { display: none; padding: 54px 20px; border: 1px dashed #bdcecb; border-radius: 7px; text-align: center; color: var(--muted); }
    .empty-state.visible { display: block; }
    .loading { padding: 20px; color: var(--muted); text-align: center; font-size: .8rem; }
    .sentinel { height: 1px; }
    @media (min-width: 641px) and (max-width: 820px) { .toolbar { grid-template-columns: minmax(0, 1fr) auto; } .stats { grid-column: 1 / -1; justify-content: space-between; } }
    @media (max-width: 640px) { .topbar-inner { padding: 18px 16px 14px; } main { padding: 23px 16px 52px; } .masthead { display: block; } .resume-button { margin-top: 14px; } .toolbar { grid-template-columns: 1fr; gap: 10px; margin-top: 17px; } .voice-control { justify-content: space-between; } .stats { justify-content: space-between; } }
    @media (prefers-reduced-motion: reduce) { *, *::before, *::after { scroll-behavior: auto !important; transition: none !important; } }
  </style>
</head>
<body>
  <header class="topbar">
    <div class="topbar-inner">
      <div class="masthead">
        <div><p class="kicker">Vocabulary studio</p><h1>考研生词</h1><p class="subtitle">按自己的节奏，逐词建立长期记忆</p></div>
        <button id="resume" class="resume-button" type="button" aria-label="回到上次学习位置"><span aria-hidden="true">↗</span> 继续上次学习</button>
      </div>
      <div class="toolbar">
        <div class="search-wrap">
          <svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" aria-hidden="true"><circle cx="11" cy="11" r="7"></circle><path d="m20 20-4-4"></path></svg>
          <input id="search" type="search" autocomplete="off" placeholder="搜索单词、释义或词性…" aria-label="搜索单词、释义或词性">
          <div id="search-results" class="search-results" role="listbox"></div>
        </div>
        <label class="voice-control" for="voice-select"><span>发音</span><select id="voice-select" aria-label="选择发音口音"><option value="en-US">美式 English (US)</option><option value="en-GB">英式 English (UK)</option></select></label>
        <div class="stats"><span><strong id="visible-count">0</strong> / <span id="total-count">0</span> 词</span><span class="meter" title="学习进度"><i id="progress-meter"></i></span><span id="progress-label">未开始</span></div>
      </div>
    </div>
  </header>
  <main>
    <div class="section-line"><span id="section-title">全部词汇</span><span id="loaded-count">正在准备…</span></div>
    <section id="grid" aria-live="polite"></section>
    <div id="empty" class="empty-state">没有找到匹配的单词，请换个关键词试试。</div>
    <div id="loading" class="loading">正在加载词卡…</div>
    <div id="sentinel" class="sentinel" aria-hidden="true"></div>
  </main>
  <script>
    const VOCAB = __VOCAB_JSON__;
    const PROGRESS_KEY = "qwentoword-progress-v1";
    const BATCH_SIZE = 36;
    const state = { query: "", matches: VOCAB, rendered: 0, lastId: null };
    const grid = document.getElementById("grid");
    const loading = document.getElementById("loading");
    const empty = document.getElementById("empty");
    const search = document.getElementById("search");
    const searchResults = document.getElementById("search-results");
    const voiceSelect = document.getElementById("voice-select");
    const speechState = { voices: [] };
    const esc = value => String(value ?? "").replace(/[&<>"']/g, character => ({"&":"&amp;","<":"&lt;",">":"&gt;","\"":"&quot;","'":"&#39;"}[character]));
    function loadProgress() { try { return JSON.parse(localStorage.getItem(PROGRESS_KEY) || "null"); } catch (_) { return null; } }
    function saveProgress(card) {
      if (!card) return;
      state.lastId = card.dataset.id;
      try { localStorage.setItem(PROGRESS_KEY, JSON.stringify({ id: state.lastId, scrollY: window.scrollY, at: Date.now() })); } catch (_) {}
      updateProgress();
      document.querySelectorAll(".card.current").forEach(node => node.classList.remove("current"));
      card.classList.add("current");
    }
    function progressId(progress) {
      if (!progress) return null;
      return progress.id && VOCAB.some(item => item.id === String(progress.id)) ? String(progress.id) : null;
    }
    function updateProgress() {
      const index = VOCAB.findIndex(item => item.id === state.lastId);
      const ratio = index < 0 ? 0 : (index + 1) / Math.max(1, VOCAB.length);
      document.getElementById("progress-meter").style.width = `${ratio * 100}%`;
      document.getElementById("progress-label").textContent = index < 0 ? "未开始" : `已学至 ${state.lastId}`;
    }
    function iconSpeaker() { return '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" aria-hidden="true"><path d="M11 5 6 9H3v6h3l5 4V5Z"></path><path d="M15.5 8.5a5 5 0 0 1 0 7M18.5 5.5a9 9 0 0 1 0 13"></path></svg>'; }
    function refreshVoices() { speechState.voices = "speechSynthesis" in window ? window.speechSynthesis.getVoices() : []; }
    function voiceFor(locale) {
      const normalisedLocale = locale.toLowerCase();
      return speechState.voices.find(voice => voice.lang.toLowerCase() === normalisedLocale)
        || speechState.voices.find(voice => voice.lang.toLowerCase().startsWith(`${normalisedLocale}-`));
    }
    function speak(word, rate) {
      if (!("speechSynthesis" in window)) { alert("当前浏览器不支持单词朗读"); return; }
      window.speechSynthesis.cancel();
      const utterance = new SpeechSynthesisUtterance(word);
      const locale = voiceSelect.value || "en-US";
      utterance.lang = locale; utterance.rate = rate || .86; utterance.pitch = 1;
      const voice = voiceFor(locale);
      if (voice) utterance.voice = voice;
      window.speechSynthesis.speak(utterance);
    }
    function openDictionary(word) {
      const term = String(word || "").trim();
      if (!term) return;
      window.open(`https://dict.youdao.com/result?word=${encodeURIComponent(term)}&lang=en`, "_blank", "noopener,noreferrer");
    }
    function getPhonetic(item) { return item.phonetic || ""; }
    function cardMarkup(item) {
      const phonetic = getPhonetic(item);
      const phoneticMarkup = phonetic
        ? `<button class="phonetic-button" type="button" data-speak-phonetic="${esc(item.word)}" aria-label="朗读 ${esc(item.word)} 的音标"><span class="phonetic">/${esc(phonetic.replace(/^\/+|\/+$/g, ""))}/</span></button>`
        : `<button class="phonetic-button" type="button" data-speak-phonetic="${esc(item.word)}" aria-label="朗读 ${esc(item.word)}"><span class="phonetic missing">音标待补充</span></button>`;
      const fields = [["词形变化", item.transformation], ["记忆方法", item.memory_techniques], ["近义辨析", item.similar_words]];
      return `<article class="card" id="word-card-${esc(item.id)}" data-id="${esc(item.id)}"><div class="card-head"><div><h2 class="word" data-word="${esc(item.word)}" tabindex="0" title="双击查询有道词典">${esc(item.word)}</h2><div class="word-meta">${phoneticMarkup}<span aria-hidden="true">·</span><span>${esc(item.part_of_speech) || "词性待补充"}</span></div></div><div class="badges"><span class="word-id">${esc(item.id)}</span><button class="speak" type="button" data-speak="${esc(item.word)}" aria-label="朗读 ${esc(item.word)}" title="朗读单词">${iconSpeaker()}</button>${item.number ? `<div class="number">${esc(item.number)}</div>` : ""}</div></div>${item.translate ? `<p class="meaning">${esc(item.translate)}</p>` : `<p class="meaning detail-value empty">释义待补充</p>`}<div class="details">${fields.map(([label, value]) => `<div class="detail"><span class="detail-label">${label}</span><span class="detail-value${value ? "" : " empty"}">${esc(value) || "待补充"}</span></div>`).join("")}</div></article>`;
    }
    function renderNextBatch() {
      if (state.rendered >= state.matches.length) { loading.textContent = state.matches.length ? "已加载全部匹配词汇" : ""; return; }
      const end = Math.min(state.rendered + BATCH_SIZE, state.matches.length);
      const fragment = document.createDocumentFragment();
      const holder = document.createElement("div");
      holder.innerHTML = state.matches.slice(state.rendered, end).map(cardMarkup).join("");
      while (holder.firstElementChild) fragment.appendChild(holder.firstElementChild);
      grid.appendChild(fragment); state.rendered = end;
      document.getElementById("loaded-count").textContent = `已显示 ${end} / ${state.matches.length}`;
      if (state.rendered >= state.matches.length) loading.textContent = state.matches.length ? "已加载全部匹配词汇" : "";
      observeCards();
      updateVisibleCount();
    }
    function observeCards() { document.querySelectorAll(".card:not([data-observed])").forEach(card => { card.dataset.observed = "1"; cardObserver.observe(card); }); }
    function updateVisibleCount() { document.getElementById("visible-count").textContent = state.matches.length; document.getElementById("total-count").textContent = VOCAB.length; document.getElementById("section-title").textContent = state.query ? `搜索 “${state.query}”` : "全部词汇"; empty.classList.toggle("visible", !state.matches.length); }
    function resetResults(matches, query) { state.matches = matches; state.query = query; state.rendered = 0; cardObserver.disconnect(); grid.innerHTML = ""; loading.textContent = matches.length ? "正在加载词卡…" : ""; updateVisibleCount(); renderNextBatch(); }
    function searchItems(query) { const normalized = query.trim().toLowerCase(); if (!normalized) return VOCAB; return VOCAB.filter(item => [item.id, item.word, item.translate, item.part_of_speech, item.memory_techniques, item.similar_words].join(" ").toLowerCase().includes(normalized)); }
    function showResults(query) {
      const matches = searchItems(query).slice(0, 8); searchResults.innerHTML = matches.map(item => `<button class="search-result" type="button" data-result-id="${esc(item.id)}" role="option"><span>${esc(item.word)}</span><small>${esc(item.translate || item.part_of_speech || "")}</small></button>`).join(""); searchResults.classList.toggle("open", Boolean(query.trim() && matches.length));
    }
    function locate(id) { const index = state.matches.findIndex(item => item.id === id); if (index < 0) return; while (state.rendered <= index) renderNextBatch(); requestAnimationFrame(() => { const target = document.getElementById(`word-card-${id}`); if (target) { target.scrollIntoView({ behavior: "smooth", block: "center" }); saveProgress(target); } }); }
    const cardObserver = new IntersectionObserver(entries => { const visible = entries.filter(entry => entry.isIntersecting && entry.intersectionRatio >= .52); if (visible.length) { visible.sort((a, b) => { const first = a.target.getBoundingClientRect(); const second = b.target.getBoundingClientRect(); return first.top - second.top || first.left - second.left; }); saveProgress(visible[0].target); } }, { threshold: [.52] });
    const sentinelObserver = new IntersectionObserver(entries => { if (entries[0].isIntersecting) renderNextBatch(); }, { rootMargin: "650px" });
    document.addEventListener("click", event => { const speakButton = event.target.closest("[data-speak]"); if (speakButton) speak(speakButton.dataset.speak); const ipaButton = event.target.closest("[data-speak-phonetic]"); if (ipaButton) speak(ipaButton.dataset.speakPhonetic, .63); const result = event.target.closest("[data-result-id]"); if (result) { const item = VOCAB.find(row => row.id === result.dataset.resultId); if (item) { search.value = item.word; resetResults([item], item.word); searchResults.classList.remove("open"); locate(item.id); } } });
    document.addEventListener("dblclick", event => { const word = event.target.closest("[data-word]"); if (word) openDictionary(word.dataset.word); });
    search.addEventListener("input", () => { const value = search.value; showResults(value); resetResults(searchItems(value), value.trim()); });
    search.addEventListener("keydown", event => { if (event.key === "Escape") { searchResults.classList.remove("open"); search.blur(); } if (event.key === "Enter") { const first = searchItems(search.value)[0]; if (first) { searchResults.classList.remove("open"); locate(first.id); } } });
    document.addEventListener("click", event => { if (!event.target.closest(".search-wrap")) searchResults.classList.remove("open"); });
    document.getElementById("resume").addEventListener("click", () => { const id = progressId(loadProgress()); if (id) { search.value = ""; resetResults(VOCAB, ""); locate(id); } });
    document.getElementById("total-count").textContent = VOCAB.length;
    refreshVoices();
    if ("speechSynthesis" in window && "onvoiceschanged" in window.speechSynthesis) window.speechSynthesis.addEventListener("voiceschanged", refreshVoices);
    updateVisibleCount(); updateProgress(); renderNextBatch(); sentinelObserver.observe(document.getElementById("sentinel"));
    const progress = loadProgress();
    const savedProgressId = progressId(progress);
    if (savedProgressId) { document.getElementById("resume").classList.add("visible"); state.lastId = savedProgressId; updateProgress(); setTimeout(() => locate(savedProgressId), 500); }
  </script>
</body>
</html>
'''


def generate(
    input_path: Path,
    output_path: Path,
    sheet_name: str | None = None,
    words_path: Path | None = None,
) -> int:
    phonetic_path = words_path or input_path.with_name("words.json")
    update_workbook(input_path, phonetic_path, sheet_name)
    rows = load_rows(input_path, sheet_name)
    serialised = (
        json.dumps(rows, ensure_ascii=False, separators=(",", ":"))
        .replace("<", "\\u003c")
        .replace("\u2028", "\\u2028")
        .replace("\u2029", "\\u2029")
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(HTML_TEMPLATE.replace("__VOCAB_JSON__", serialised), encoding="utf-8")
    return len(rows)


def main() -> None:
    args = parse_args()
    output = args.output or args.input.with_suffix(".html")
    words_path = args.words or args.input.with_name("words.json")
    count = generate(args.input, output, args.sheet, words_path)
    print(f"Generated {output} with {count} vocabulary cards.")


if __name__ == "__main__":
    main()
